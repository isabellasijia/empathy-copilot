from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .db import database, initialize


TICKET_SHEETS = (
    "补发换货工单",
    "线下打款工单",
    "物流工单",
    "不良反应工单",
    "售后退货工单",
)


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _clean_datetime(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    text = _clean_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return text


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ")
    return value


def _records(sheet: Any) -> Iterable[tuple[int, dict[str, Any]]]:
    rows = sheet.iter_rows(values_only=True)
    headers = [_clean_text(value) or f"column_{index}" for index, value in enumerate(next(rows))]
    for row_number, row in enumerate(rows, start=2):
        record = {headers[index]: _json_value(value) for index, value in enumerate(row)}
        if any(value not in (None, "") for value in record.values()):
            yield row_number, record


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _first(record: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if record.get(key) not in (None, ""):
            return record[key]
    return None


def _ticket_values(sheet_name: str, row_number: int, record: dict[str, Any]) -> tuple[Any, ...]:
    ticket_id = _clean_text(record.get("工单号"))
    category = _first(record, "工单类型", "打款类型", "问题类型", "类型", "包裹类型")
    reason = _first(record, "售后原因", "退款问题类型", "问题类型", "症状描述", "退货原因")
    status = _first(record, "工单状态", "任务状态", "转账状态")
    return (
        ticket_id,
        _clean_text(record.get("会话ID")),
        _clean_text(record.get("关联订单号")),
        _clean_text(record.get("买家昵称")),
        _clean_text(record.get("店铺")),
        sheet_name.removesuffix("工单"),
        _clean_text(category),
        _clean_text(reason),
        _clean_text(status),
        _clean_text(record.get("处理人")),
        _clean_datetime(record.get("创建时间")),
        _clean_datetime(record.get("完成时间")),
        _clean_text(record.get("发出商品货号")),
        _clean_text(_first(record, "发出商品名称", "使用商品")),
        _clean_text(
            _first(
                record,
                "补发物流单号",
                "相关物流单号",
                "问题包裹物流单号",
                "退货物流单号",
            )
        ),
        _clean_text(record.get("产品批次号")),
        sheet_name,
        row_number,
        json.dumps(record, ensure_ascii=False, default=str),
    )


def import_workbook(workbook_path: Path, database_path: Path, force: bool = False) -> dict[str, Any]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到官方数据文件: {workbook_path}")

    initialize(database_path)
    workbook_hash = _hash_file(workbook_path)
    with database(database_path) as connection:
        existing = connection.execute(
            "SELECT value FROM app_meta WHERE key = 'workbook_sha256'"
        ).fetchone()
        if existing and existing["value"] == workbook_hash and not force:
            report = connection.execute(
                "SELECT value FROM app_meta WHERE key = 'etl_report'"
            ).fetchone()
            return json.loads(report["value"]) if report else {"status": "unchanged"}

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet_names = set(workbook.sheetnames)
    required = {"聊天记录", "订单", *TICKET_SHEETS}
    missing_sheets = sorted(required - sheet_names)
    if missing_sheets:
        raise ValueError(f"官方数据缺少工作表: {', '.join(missing_sheets)}")

    with database(database_path) as connection:
        for table in (
            "action_log",
            "risk_events",
            "commitments",
            "analysis_cache",
            "tickets",
            "orders",
            "messages",
            "conversations",
            "raw_records",
        ):
            connection.execute(f"DELETE FROM {table}")

        chat_records = list(_records(workbook["聊天记录"]))
        grouped: dict[str, list[tuple[int, dict[str, Any]]]] = defaultdict(list)
        for row_number, record in chat_records:
            session_id = _clean_text(record.get("会话ID"))
            message_id = _clean_text(record.get("message_id"))
            if not session_id or not message_id:
                continue
            grouped[session_id].append((row_number, record))
            connection.execute(
                "INSERT INTO raw_records(source_sheet, source_row, record_key, payload_json) VALUES (?, ?, ?, ?)",
                ("聊天记录", row_number, message_id, json.dumps(record, ensure_ascii=False, default=str)),
            )

        for session_id, items in grouped.items():
            ordered = sorted(items, key=lambda item: int(item[1].get("消息序号") or 0))
            first = ordered[0][1]
            times = [_clean_datetime(item[1].get("发送时间")) for item in ordered]
            times = [value for value in times if value]
            connection.execute(
                """
                INSERT INTO conversations(
                    session_id, buyer_nickname, shop, scene_major, scene_minor,
                    started_at, last_message_at, message_count
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    session_id,
                    _clean_text(first.get("买家昵称")) or "未知客户",
                    _clean_text(first.get("店铺")),
                    _clean_text(first.get("scene_major")),
                    _clean_text(first.get("scene_minor")),
                    min(times) if times else None,
                    max(times) if times else None,
                    len(ordered),
                ),
            )

        for row_number, record in chat_records:
            session_id = _clean_text(record.get("会话ID"))
            message_id = _clean_text(record.get("message_id"))
            if not session_id or not message_id or session_id not in grouped:
                continue
            content_type = "image" if _clean_text(record.get("内容类型")) == "图片" else "text"
            image_path = _clean_text(record.get("image_path"))
            resolved_image = workbook_path.parent / image_path if image_path else None
            connection.execute(
                """
                INSERT INTO messages(
                    message_id, session_id, message_seq, sent_at, role, sender, text,
                    content_type, image_path, image_available, order_id, ticket_id, source_row
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    message_id,
                    session_id,
                    int(record.get("消息序号") or 0),
                    _clean_datetime(record.get("发送时间")),
                    "customer" if _clean_text(record.get("角色")) == "买家" else "agent",
                    _clean_text(record.get("发送方")),
                    _clean_text(record.get("message_text")) or "",
                    content_type,
                    image_path,
                    int(bool(resolved_image and resolved_image.exists())),
                    _clean_text(record.get("关联订单号")),
                    _clean_text(record.get("关联工单号")),
                    row_number,
                ),
            )

        order_count = 0
        for row_number, record in _records(workbook["订单"]):
            order_id = _clean_text(record.get("订单号"))
            if not order_id:
                continue
            order_count += 1
            connection.execute(
                "INSERT INTO raw_records(source_sheet, source_row, record_key, payload_json) VALUES (?, ?, ?, ?)",
                ("订单", row_number, order_id, json.dumps(record, ensure_ascii=False, default=str)),
            )
            connection.execute(
                """
                INSERT INTO orders(
                    order_id, session_id, buyer_nickname, shop, sku, product_name,
                    quantity, unit_price, paid_amount, status, ordered_at, paid_at,
                    shipped_at, carrier, tracking_no, province, city, gift, buyer_note,
                    source_row, raw_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    order_id,
                    _clean_text(record.get("会话ID")),
                    _clean_text(record.get("买家昵称")),
                    _clean_text(record.get("店铺")),
                    _clean_text(record.get("商品货号")),
                    _clean_text(record.get("商品名称")),
                    record.get("数量"),
                    record.get("单价(元)"),
                    record.get("实付金额(元)"),
                    _clean_text(record.get("订单状态")),
                    _clean_datetime(record.get("下单时间")),
                    _clean_datetime(record.get("付款时间")),
                    _clean_datetime(record.get("发货时间")),
                    _clean_text(record.get("快递公司")),
                    _clean_text(record.get("物流单号")),
                    _clean_text(record.get("收货省")),
                    _clean_text(record.get("收货市")),
                    _clean_text(record.get("赠品")),
                    _clean_text(record.get("买家留言")),
                    row_number,
                    json.dumps(record, ensure_ascii=False, default=str),
                ),
            )

        ticket_count = 0
        for sheet_name in TICKET_SHEETS:
            for row_number, record in _records(workbook[sheet_name]):
                ticket_id = _clean_text(record.get("工单号"))
                if not ticket_id:
                    continue
                ticket_count += 1
                connection.execute(
                    "INSERT INTO raw_records(source_sheet, source_row, record_key, payload_json) VALUES (?, ?, ?, ?)",
                    (sheet_name, row_number, ticket_id, json.dumps(record, ensure_ascii=False, default=str)),
                )
                connection.execute(
                    """
                    INSERT INTO tickets(
                        ticket_id, session_id, order_id, buyer_nickname, shop, ticket_kind,
                        category, reason, status, owner, created_at, completed_at, product_sku,
                        product_name, tracking_no, batch_no, source_sheet, source_row, raw_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    _ticket_values(sheet_name, row_number, record),
                )

        image_rows = connection.execute(
            "SELECT COUNT(*) AS total, SUM(image_available) AS available FROM messages WHERE content_type = 'image'"
        ).fetchone()
        report = {
            "status": "imported",
            "workbook": workbook_path.name,
            "conversations": len(grouped),
            "messages": len(chat_records),
            "orders": order_count,
            "tickets": ticket_count,
            "image_messages": image_rows["total"],
            "available_images": image_rows["available"] or 0,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
        }
        connection.execute(
            "INSERT OR REPLACE INTO app_meta(key, value) VALUES ('workbook_sha256', ?)",
            (workbook_hash,),
        )
        connection.execute(
            "INSERT OR REPLACE INTO app_meta(key, value) VALUES ('etl_report', ?)",
            (json.dumps(report, ensure_ascii=False),),
        )

    workbook.close()
    return report

